from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt


def gerar_relatorio_alunos(serie, alunos):
    documento = Document()

    # Título
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = titulo.add_run("EduTrack")
    run.bold = True
    run.font.size = Pt(20)

    # Subtítulo
    subtitulo = documento.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitulo.add_run("Relatório de Alunos")
    run.bold = True
    run.font.size = Pt(14)

    # Série
    paragrafo_serie = documento.add_paragraph()
    run = paragrafo_serie.add_run(f"Série: {serie}")
    run.bold = True

    # Tabela
    tabela = documento.add_table(
        rows=1,
        cols=4
    )

    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = "Table Grid"

    cabecalho = tabela.rows[0].cells

    cabecalho[0].text = "Nome"
    cabecalho[1].text = "Notas"
    cabecalho[2].text = "Média"
    cabecalho[3].text = "Status"

    for aluno in alunos:
        linha = tabela.add_row().cells

        linha[0].text = aluno.nome
        linha[1].text = ", ".join(
            str(nota.valor) for nota in aluno.notas
        )
        linha[2].text = str(round(aluno.media, 2))

        status = (
            "Aprovado"
            if aluno.media >= 5
            else "Reprovado"
        )

        linha[3].text = status

    # Salvar o documento em memória
    arquivo = BytesIO()

    documento.save(arquivo)

    arquivo.seek(0)

    return arquivo
